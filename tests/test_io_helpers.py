from docx import Document

from trapo_app import io_helpers


def _table_doc(path, rows):
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.save(path)


def _first_column(path):
    return [row.cells[0].text for row in Document(path).tables[0].rows]


def test_sort_word_table_orders_rows_and_keeps_header(tmp_path):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    _table_doc(src, [["Treffpunkt", "Name"], ["B", "1"], ["A", "2"], ["C", "3"], ["B", "4"]])
    io_helpers.sort_word_table(str(src), str(dst), sort_column=0, sort_order=["A", "B"])
    assert _first_column(dst) == ["Treffpunkt", "A", "B", "B", "C"]


def test_read_docx_uses_first_row_as_header(tmp_path):
    path = tmp_path / "t.docx"
    _table_doc(path, [["Name", "Kontakt"], ["Rex", "x"]])
    df = io_helpers.read_docx(str(path))
    assert list(df.columns) == ["Name", "Kontakt"]
    assert df["Name"].tolist() == ["Rex"]


def test_filter_stopps_returns_each_file_once():
    files = ["a-SÜDWEST-V1.docx", "b-NORD.docx", "c-other.docx"]
    assert io_helpers.filter_stopps(files) == ["a-SÜDWEST-V1.docx", "b-NORD.docx"]


def test_split_word_table_writes_one_file_per_part(tmp_path):
    src = tmp_path / "trapo.docx"
    _table_doc(src, [["Treffpunkt", "Name"], ["B", "1"], ["A", "2"], ["C", "3"], ["B", "4"]])
    parts = {"Nord": ["A", "C"], "Süd/West": ["B"], "Leer": ["X"]}
    written = io_helpers.split_word_table(str(src), str(tmp_path / "trapo"), 0, parts)
    assert set(written) == {"Nord", "Süd/West"}  # parts without rows are not written
    assert written["Nord"] == (str(tmp_path / "trapo_Nord.docx"), 2)
    assert _first_column(tmp_path / "trapo_Nord.docx") == ["Treffpunkt", "A", "C"]
    assert _first_column(tmp_path / "trapo_Süd-West.docx") == ["Treffpunkt", "B", "B"]
