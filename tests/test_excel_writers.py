import io

import openpyxl
import pandas as pd
from PIL import Image

from trapo_app import excel_writers, table_helpers


def _png(size):
    buffer = io.BytesIO()
    Image.new("RGB", size, "red").save(buffer, "PNG")
    return buffer.getvalue()


def test_write_df_to_excel_roundtrip(tmp_path):
    path = tmp_path / "out.xlsx"
    df = pd.DataFrame({"Name": ["Rex", "Bello"], "Chip": ["1", "22"]})
    excel_writers.write_df_to_excel(df, path, "Blatt")
    ws = openpyxl.load_workbook(path)["Blatt"]
    assert [[c.value for c in row] for row in ws.iter_rows()] == [["Name", "Chip"], ["Rex", "1"], ["Bello", "22"]]


def test_distance_sheet_layout(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    df = pd.DataFrame({
        "Photo": ["img_0", "img_1", "img_2"], "Name": ["Rex", "Bello", "Luna"],
        "Treffpunkt": ["Nord", "Nord", "Süd"], "Entfernung": [10, 20, 30],
    })
    images = {"img_0": _png((300, 200)), "img_1": _png((50, 40)), "img_2": _png((120, 400))}
    excel_writers.save_distance_sheets(["chat.docx"], [table_helpers.insert_headers(df)], [images])

    ws = openpyxl.load_workbook(tmp_path / "chat_Entfernung.xlsx").active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows == [
        ["Nord", None, None, None, None],
        ["Nr.", "Photo", "Name", "Treffpunkt", "Entfernung"],
        [1, None, "Rex", "Nord", 10],
        [2, None, "Bello", "Nord", 20],
        ["Süd", None, None, None, None],
        ["Nr.", "Photo", "Name", "Treffpunkt", "Entfernung"],
        [1, None, "Luna", "Süd", 30],
    ]
    assert sorted(map(str, ws.merged_cells.ranges)) == ["A1:E1", "A5:E5"]
    assert len(ws._images) == 3
    # the picture of the first animal sits in its data row (Excel row 3)
    assert ws._images[0].anchor._from.row == 2


def test_ro_sheet_has_title_and_footer(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    df = pd.DataFrame(columns=table_helpers.RO_HEADERS, data=[["1"] + [None] * 16])
    excel_writers.save_ro_excel([df], ["liste.docx"])
    ws = openpyxl.load_workbook(tmp_path / "liste.xlsx").active
    assert ws["A1"].value == excel_writers.RO_TITLE
    assert ws["A4"].value == excel_writers.RO_FOOTER_LEFT
    assert ws.cell(4, len(table_helpers.RO_HEADERS) - 2).value == excel_writers.RO_FOOTER_RIGHT
